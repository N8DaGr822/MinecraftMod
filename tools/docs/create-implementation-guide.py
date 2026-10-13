"""Build the shareable implementation guide from current source catalogs and reviewed feature notes."""
from pathlib import Path
import json
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'output' / 'documents'
OUT.mkdir(parents=True, exist_ok=True)
README = (ROOT / 'README.md').read_text(encoding='utf-8')
LANG = json.loads((ROOT / 'src/main/resources/assets/darkspawn/lang/en_us.json').read_text(encoding='utf-8'))
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.65)
sec.left_margin = sec.right_margin = Inches(.7)
sec.footer_distance = Inches(.28)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption']:
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_after = Pt(6)
    for border in st.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(32)
doc.styles['Heading 1'].font.size = Pt(22)
doc.styles['Heading 2'].font.size = Pt(14)
doc.styles['Heading 3'].font.size = Pt(11)
doc.styles['Caption'].font.size = Pt(9)
doc.core_properties.title = 'Darkspawn Implementation Guide'
doc.core_properties.subject = 'Implemented gameplay features and illustrated catalogs'
doc.core_properties.author = 'Darkspawn'
doc.core_properties.keywords = 'Minecraft, Fabric, Darkspawn, player guide'
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Darkspawn  |  October 13 2026  |  ').font.size = Pt(8)
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); footer._p.append(fld)

def clean(text):
    text = re.sub(r'\[([^]]+)\]\([^)]*\)', r'\1', str(text))
    return text.replace('**', '').replace('`', '').replace('—', ' - ').replace('–', '-').replace('’', "'")

def p(text, style=None):
    return doc.add_paragraph(clean(text), style)

def h(text, level=2):
    return doc.add_heading(text, level)

def chapter(title):
    heading = doc.add_heading(title, 1)
    heading.paragraph_format.page_break_before = True

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = widths or [7.1 / len(headers)] * len(headers)
    for col, width in zip(t.columns, widths): col.width = Inches(width)
    for i, text in enumerate(headers): t.rows[0].cells[i].text = clean(text)
    for row in rows:
        cells = t.add_row().cells
        for i, text in enumerate(row): cells[i].text = clean(text)
    borders = OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
    t._tbl.tblPr.append(borders)
    for ri,row in enumerate(t.rows):
        cant=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(cant)
        if ri==0:
            repeat=OxmlElement('w:tblHeader'); row._tr.get_or_add_trPr().append(repeat)
        for ci,cell in enumerate(row.cells):
            cell.width=Inches(widths[ci]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'243E50' if ri==0 else ('F2F5F7' if ri%2 else 'FFFFFF')); tcPr.append(shade)
            margin=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge); e.set(qn('w:w'),'60'); e.set(qn('w:type'),'dxa'); margin.append(e)
            tcPr.append(margin)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(1); para.paragraph_format.space_before=Pt(1)
                para.paragraph_format.line_spacing=1.03
                if ri==0: para.paragraph_format.keep_with_next=True
                if widths[ci]<.85: para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.font.size=Pt(9.5)
                    if ri==0: run.bold=True; run.font.color.rgb=RGBColor(255,255,255)
    p('')
    return t

def source_table(marker):
    lines=README[README.index(marker):].splitlines()
    rows=[]
    for line in lines:
        if line.startswith('|'):
            cells=[clean(c.strip()) for c in line.strip('|').split('|')]
            if not all(re.fullmatch(r'[:\- ]+', c) for c in cells):rows.append(cells)
        elif rows:break
    return rows[0], rows[1:]

def use_table(marker, widths):
    headers,rows=source_table(marker);return table(headers,rows,widths)

def figure(path, caption, width=6.8, max_height=5.7):
    from PIL import Image
    path=ROOT/path
    with Image.open(path) as im: width=min(width,max_height*im.width/im.height)
    para=doc.add_paragraph(); para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.keep_with_next=True
    pic=para.add_run().add_picture(str(path),width=Inches(width))
    pic._inline.docPr.set('descr',caption)
    p(caption,'Caption')

def item_name(id):
    ns,name=id.split(':',1) if ':' in id else ('darkspawn',id)
    return LANG.get(f'item.{ns}.{name}',LANG.get(f'block.{ns}.{name}',name.replace('_',' ').title()))

p('Darkspawn Implementation Guide','Title')
p('Gameplay features and illustrated catalogs','Subtitle')
p('Current development implementation  |  October 13 2026')
figure('art/source/project-icon.png','Darkspawn tree spirit emblem.',width=2.25,max_height=2.25)
p('Darkspawn expands Minecraft with biome encounters, rideable creatures, cooking, equipment progression and shared exploration systems. This guide explains what players can currently use, where to find it, and how the systems work together.')
p('This edition describes the development source for Minecraft 26.3 and Fabric, including the new mount habitats and five additional mounts. It includes changes made after the last packaged JAR. An older download may not contain every feature shown here.')
p('The mod remains in alpha. Implemented features and playable prototypes are identified throughout. Planned ideas are listed separately at the end. Asset previews show the actual project models; they are not gameplay screenshots.')
h('Contents')
for text in ['Starting and progression','Mounts and natural habitats','Boss encounters','Boss rewards and equipment powers','Regional creatures','Cooking and agriculture','Armor and climate','Inventory and storage','Bonfires and exploration structures','OneBlock and Forsaken','Darkspawn Journal','Presentation and prototype boundaries','Cooking Station recipe catalog']:
    p(text)

chapter('Starting and progression')
p('Install matching Darkspawn versions on every client and the Fabric server, together with Fabric API and GeckoLib. The current source targets Minecraft 26.3, Fabric Loader 0.19.5, Fabric API 0.162.0+26.3 and GeckoLib 5.5.7. Restart after updating so dimension and biome registrations load.')
table(['Available before the Ender Dragon','Unlocked by the world dragon victory'],[
('Inventory sorting and upgrades; armor loadouts; storage networks; ordinary cooking, farming and Chef trades; armor styling; expedition gear and climate enchantments; all mount habitats; bonfires; shared OneBlock.','Boss rituals, natural regional creature populations, natural Endborn, regional atmosphere, and activation of Forsaken gates.'),
('Druid Shrines, Hunter Camps and research bunkers can be explored before victory.','Boss rewards, legendary cores, premium ingredients and later progression become obtainable through the unlocked encounters.')],[3.55,3.55])
p('Dragon progression is shared by the world. Permanent hearts, discoveries, inventory capacity and armor-loadout unlocks belong to individual players. Ordinary items and crafted equipment can be shared. Tamed mounts belong to their current owner; only that player can ride them.')
h('Implemented scope')
table(['System','Current contents'],[
('Travel','14 rideable mount types with natural habitats; flight, gliding, grounded travel and swimming.'),
('Encounters','16 boss families, 82 regional creatures and discoveries, Endborn, and 16 encounter-minion types.'),
('Food','59 Cooking Station recipes: 50 prepared meals, six feasts and three staple recipes. Nine crop types including regional variants.'),
('Equipment','15 standard armor styles, five legendary sets, five expedition families, 16 melee empowerments and ten shield, bow or armor powers.'),
('Progression','Personal boss-heart upgrades, recipe discoveries, dietary variety, meal collection and advancements.'),
('Building and storage','Sorting Tubes, Dump Chests, Sorting Chests, Bank Chests, a Barracks Bonfire and discovery landmarks.'),
('Dimensions','Shared OneBlock progression and Forsaken city expeditions with infected threats and equipment rewards.')],[1.3,5.8])
h('Optional builder supplies')
p('The Dragon Starter Chest is a Creative or command-only equipment kit. It contains enchanted diamond gear, a bow, shield, food, building supplies and potions for the vanilla dragon fight. It has no Survival recipe or automatic world placement. Breaking the placed chest does not regenerate another kit.')
p('The Testing Camp item builds a 21 by 21 supply camp on suitable clear, flat ground. It includes crafting and cooking, enchanting, crops, an inactive bonfire and labeled supplies. Placement refuses obstructions rather than partially overwriting a build. These helpers are optional; they are not normal progression rewards.')

chapter('Mounts and natural habitats')
figure('build/mount-contact-sheet.png','Mount asset preview. Top row: Woolly Rhino, Flying Squirrel, Poptart Cat, Riding Dragon. Second row: Winged Unicorn, Riding Golem, Sasquatch, Cerberus. Third row: Jungle Leopard, Giant Seahorse, Giant Scorpion, Swamp Crocodile. Last row: Giant Mole.',max_height=3.7)
p('Mounts are passive travel prototypes available before the dragon victory and in Peaceful. Feed the preferred item three times, at least two seconds apart, to tame a mount. Only the owner can ride, by empty-hand right-click; no saddle is required. Taming makes a wild mount persistent, and ownership survives saves and reconnects.')
table(['Control','Result'],[('WASD and mouse','Move relative to the rider facing direction.'),('Space','Jump on land; rise on flying or swimming mounts.'),('X','Descend in flight, dive while swimming, or descend faster while gliding.'),('V plus movement','Dig a tunnel while riding a Giant Mole.'),('Shift','Dismount. A nearby safe floor is preferred; emergency midair exits grant Slow Falling.'),('Release movement','Flying mounts brake into a hover. Swimmers settle in water. The squirrel continues descending.')],[1.5,5.6])
h('Movement boundaries')
p('Leopards have a higher grounded jump and a larger safe-fall allowance; they cannot climb walls or fly. Seahorses and crocodiles can move up and down in water without drowning themselves. Their riders still need air, Water Breathing or appropriate diving equipment. The squirrel needs height to glide and cannot gain altitude. Ground mounts retain gravity and fall damage.')
p('Flyers on Stay descend, except unclaimed sky cats, which hover. Follow steers toward the owner. Losing control packets stops active steering. Mounting and movement check rider headroom. Companion care and combat are described below.')

h('Where mounts appear')
table(['Mount','Natural habitat','Travel'],[
('Griffin','Exposed jagged, frozen or stony peaks at Y128 or higher.','Flight and hover'),
('Cerberus','Nether biomes below Y120 on dry supported ground.','Fast ground travel'),
('Flying Squirrel','Surface forests and taigas.','Ground travel and downward gliding'),
('Winged Unicorn','End island surfaces at Y30 or higher.','Flight and hover'),
('Woolly Rhino','Snowy plains, ice spikes, snowy taiga, grove, snowy slopes, frozen or jagged peaks.','Ground travel'),
('Riding Dragon','Covered Overworld caves below Y48, with lava within six blocks horizontally and three vertically.','Flight and hover'),
('Riding Golem','Covered Overworld caves below Y48.','Slow ground travel'),
('Sasquatch','Snowy taiga and grove surfaces.','Ground travel'),
('Poptart Cat','Very rare open upper Overworld air, normally Y272-311.','Flight, hover and rear rainbow trail'),
('Jungle Leopard','Jungle, sparse jungle and bamboo jungle surfaces.','Fast ground travel and a higher jump'),
('Giant Seahorse','Ocean water with room for the full body and water above and below.','Controlled swimming; very slow on land'),
('Giant Scorpion','Desert surfaces.','Ground travel'),
('Swamp Crocodile','Swamp and mangrove swamp banks near water.','Ground travel and controlled swimming'),
('Giant Mole','Super rare in Soul Sand Valley below Y120, on dry supported ground.','Ground travel and tunneling')],[1.25,4.25,1.6])
p('Ordinary mounts use native biome spawning, population caps and single-animal groups. They require collision-free space. Land encounters require supported dry ground; cave mounts need overhead terrain at least eight blocks above the spawn level. A living mount of the same species within 96 blocks on each axis blocks a nearby duplicate. Unclaimed mounts may despawn with distance. These encounters can occur in existing terrain after updating; they do not require new chunks.')
p('The sky cat keeps its separate very-rare rule: once per minute, one world-wide 1-in-20 roll is possible while a living non-spectator player is within 80 blocks of the Overworld ceiling. A successful roll tries one loaded position 32-64 blocks away horizontally. Blocked air, unloaded chunks or another cat within 128 blocks cancel the attempt. It can appear day or night. Twenty eligible minutes is the average time per successful roll, not a guaranteed spawn interval.')
p('The Giant Mole has spawn weight 1 plus a 1-in-100 roll for each otherwise-valid placement. Hold V and move to clear a tunnel with space for the rider. It preserves the floor and drops normal diamond-tier mining loot without Fortune or Silk Touch. Up to 12 blocks are cleared per five ticks. Unbreakable blocks, fluids and containers are skipped; player build restrictions, spawn protection and Fabric block-break vetoes apply. Creative digging produces no drops. It cannot dig vertically or safely cross lava.')
h('Taming and care').paragraph_format.page_break_before = True
p('Trust messages and green particles show feeding progress; hearts confirm taming. Each successful Survival feeding consumes one item. Creative feeding does not consume items. The active feeder reserves the attempt for sixty seconds after each successful feeding. Other players cannot finish it during that time. Afterward, another feeder starts from zero; the original feeder retains saved progress until someone else starts.')
table(['Mount','Preferred bonding item'],[('Griffin', 'Raw rabbit'), ('Woolly Rhino', 'Wheat'), ('Flying Squirrel', 'Sweet berries'), ('Poptart Cat', 'Cookie'), ('Riding Dragon', 'Cooked beef'), ('Winged Unicorn', 'Golden apple'), ('Riding Golem', 'Iron ingot'), ('Sasquatch', 'Apple'), ('Cerberus', 'Bone'), ('Jungle Leopard', 'Raw chicken'), ('Giant Seahorse', 'Kelp'), ('Giant Scorpion', 'Spider eye'), ('Swamp Crocodile', 'Raw cod'), ('Giant Mole', 'Beetroot')],[3.1,4.0])
p('Once tamed, only the owner can feed or ride. Each feeding heals four health points, or two hearts, with the same two-second cooldown. Full-health mounts, cooldown clicks and rejected attempts consume nothing. Iron ingots and bones act as bonding or repair items for the golem and Cerberus.')
p('Existing ownerless mounts keep their identity, name, health and position but require normal taming. There is no free first-rider claim. Previously persistent mounts stay persistent; incomplete trust does not prevent a natural mount from despawning. Tamed sky cats settle on Stay. Owned mounts take normal damage but enter veterinary recovery when defeated.')


h('Companion lifecycle').paragraph_format.page_break_before = True
h('Companion commands and cargo')
p('Sneak-right-click your companion empty-handed to open its panel within eight blocks. Follow approaches the owner within 32 blocks in the same dimension. Ground mounts navigate, flyers steer through open air and aquatic mounts swim. Seahorses stop chasing an owner on land. Following does not teleport, load chunks or dig, and terrain can block it. Stay stops following and cancels breeding readiness; flyers settle.')
p('All fourteen species have nine cargo slots. Cargo survives saves, capsules, recovery and ownership transfer. Capsules, bundles, shulker boxes and items carrying container contents cannot go inside mount cargo.')
h('Capsules and veterinary recovery')
p('Use an empty Companion Capsule on your owned, unoccupied, unleashed mount to store its identity, name, health and cargo. Use the filled capsule on safe ground to release in Stay mode. Seahorses require water around their full body. Ridden mounts, breeding requests and pending transfers block capture. Ordinary lost capsules must be recovered; veterinary reissue is limited to defeated companions.')
p('Lethal damage to an owned mount creates a recovery token. Carry a Veterinary Kit and use the token on a safe surface to restore full health, consuming one kit on success. Use a kit directly on an injured companion to heal it fully. Use a kit in air to reclaim one missing recovery token; joining also returns one. Offline recoveries remain saved. Copied tokens cannot redeem a companion twice. OneBlock companions and cargo cannot cross into normal worlds through capsules, or vice versa.')
h('Breeding and ownership transfer')
p('Set two healthy adults of the same species and owner to Stay near each other. Choose Breed for each, spending one preferred item from inventory per parent. Readiness lasts thirty seconds. A successful pairing creates one owned baby and starts a five-minute cooldown for both parents. Babies grow after twenty minutes of loaded entity ticks and cannot be ridden before then. Growth pauses while stored. Saves retain growth and cooldowns but cancel incomplete breeding requests.')
p('Select a living nearby player and choose Offer transfer. That recipient must sneak-right-click the mount empty-handed and accept within sixty seconds while both players remain within eight blocks. The offer includes cargo. Pending offers block riding, storage and breeding. Cancel or reload clears the offer. Acceptance revokes the old owner immediately and selects Stay.')
table(['Crafting table item','Ingredients and arrangement'],[
('Companion Capsule','Top: iron in center. Middle: amethyst shard, ender pearl, amethyst shard. Bottom: iron in center.'),
('Veterinary Kit','Top: white wool in center. Middle: apple, gold ingot, apple. Bottom: iron in center.')],[1.4,5.7])

h('Mount combat abilities').paragraph_format.page_break_before = True
p('Hold Mount: Ability, G by default and rebindable in Controls, to attack whenever the cooldown permits. Rhino charge also requires forward movement on dry ground. The owning rider controls each activation; mounts do not attack autonomously. An on-screen message shows the cooldown after use.')
table(['Mount','Ability','Cooldown'],[
('Riding Dragon','Aimed eight-block cone deals 8 fire damage and three seconds of burning. Fire-immune enemies retain their immunity.','4 seconds'),
('Woolly Rhino','Up to one second of faster ground travel; contact deals 10 damage and knockback, once per enemy per charge. Actual movement is required.','5 seconds'),
('Giant Scorpion','Nearest hostile in a forward cone within 3.5 blocks takes 4 damage and Poison I for five seconds. Poison immunity still applies.','3 seconds'),
('Swamp Crocodile','Nearest hostile in a forward cone within 3.5 blocks takes 8 damage and Slowness I for two seconds, on land or underwater.','2 seconds'),
('Cerberus','The mount and mounted owner are immune to fire and lava damage. Rider protection ends immediately on dismount.','Passive')],[1.2,4.85,1.05])
p('Damage is measured in health points before armor and resistance; two points equal one heart. Attacks spare players, owned companions, allies and peaceful regional wildlife. Neutral vanilla enemies must already be targeting the rider or mount. Walls block attacks, and dragon breath never ignites or destroys terrain.')
p('Misses still consume the cooldown. Cooldowns survive remounting and saving; active charges do not resume after loading. Releasing forward, stale controls, dismounting, leaving the ground, water or a wall ends a rhino charge. Server-synchronized jaw, head, wing, tail and claw poses accompany the existing sounds and particles. Attacks blend back into movement; cancellation and reloading clear active poses.')
p('Cerberus protection leaves existing potion effects untouched. It grants neither lava walking nor protection against other damage. A dismounted rider is vulnerable immediately, so dismount onto safe ground.')

chapter('Boss encounters')
p('All 16 boss families have summon rituals, models, animations, phased encounters and rewards. Bosses are summoned by players; the surrounding ecosystem creatures do not automatically start a boss fight. Every ritual needs the world dragon victory, the correct habitat, non-Peaceful difficulty and enough clear space. A nearby living boss of the same family prevents a duplicate. Failed rituals preserve the summon item.')
p('Boss health increases by half the solo value for each additional nearby Survival player, up to four players, and is fixed at summon time. Phases advance at 70% and 35% health and do not reverse through healing. Attacks use visible warnings and recovery windows. Most attack effects preserve player terrain.')
headers,rows=source_table('| Boss / entity ID | Solo HP |')
rows=[[re.sub(r' / .*','',r[0]),r[1],r[2]] for r in rows]
rows=[['Ancient Tree Spirit','800','Moss block in a supported forest clearing.'],['Giant Mutated Wolf','900','Bone block in a supported taiga at night.']]+rows
table(['Boss','Solo HP','Ritual location'],rows,[1.85,.7,4.55])
h('Summoning and arena preparation')
p("Most sigils use four outer materials, four inner materials and dragon's breath. The Cookbook is for food; boss item tooltips and the crafting recipe book explain the summon ingredients. Regional discovery materials offer an alternative summon recipe. Void Eye uses the Watcher Lens, which includes Endborn Shards. Kraken uses Abyssal Bait cast into deep open ocean water.")
p('The Tree Spirit requires a 21 by 21 clearing and 28 blocks of height. It awakens as a towering rooted tree, then sheds canopy sections as phases advance. Destroy healing saplings and use the exposed heartwood recovery window. The wolf uses chase pressure, warned pounces, howls and temporary pack support.')
h('Distinct encounter mechanics')
use_table('| Boss | Phase one |', [1.2,1.85,2.0,2.05])
p('Bone guardians, royal guardians and Kraken tentacles protect their owners with damage reduction. Destroy these supports to reopen damage opportunities. Healing nodes and turrets have limited lives and their own warnings. Encounter waves have finite budgets and do not farm unlimited rewards.')
p('Soulbound Colossus temporarily steals maximum-health hearts; defeating soul cages, waiting out the effect or ending the encounter returns them. Permanent heart progress is never deleted. Void Eye displacement checks for safe supported destinations and grants Slow Falling instead of deliberately dropping players into the void.')
p('Encounters expire without rewards when living Survival players remain away for the timeout. Boss state, phase and participants survive saves. Reloaded attacks receive a new warning rather than silently completing. Solo and multiplayer balance remain alpha values.')

chapter('Boss rewards and equipment powers')
figure('art/items/boss-essences-preview.png','Boss essence artwork preview. Essences are separate from permanent-health hearts and cooking ingredients.',max_height=3.1)
h('Permanent hearts and repeat rewards')
p('Right-click a boss heart to absorb its family identity. The first ten distinct hearts each add one normal heart, bringing the normal total to twenty. Identities eleven through fifteen add five golden absorption hearts of capacity. Duplicates are rejected without consumption. There are sixteen families, so players can choose which fifteen finish the track.')
p('Golden capacity refills after sixty seconds without damage. Joining, respawning or taking damage restarts the delay, and stronger golden-apple absorption is preserved. Heart progress survives death, reconnects and server restarts. Normal added hearts heal through normal regeneration.')
p('Nearby qualifying contributors receive individually reserved progression rewards. Each boss also has a separate premium culinary ingredient and decorative trophy, obtainable on repeat wins. The smithing essence is a shared kill drop. Helpers and abandoned encounters do not grant these rewards.')
h('Melee empowerments')
p('Smith an amethyst shard, a melee weapon and the listed boss material. The item keeps its identity, damage, name, enchantments and unrelated components. It holds one empowerment: a different power replaces the previous one, while reapplying the same power is rejected. Powers trigger on successful direct melee hits.')
use_table('| Boss drop | Weapon power |',[1.55,5.55])
h('Shield bow and armor powers')
headers,rows=source_table('| Equipment | Material | Power |')
rows += [['Shield','Ice Wyrm essence','Frostguard weakens an attacker for three seconds after a successful block.'],['Bow','Titan Boa essence','Venomshot applies Poison I for three seconds.'],['Armor','Mountain Titan essence','Highland reduces the fall-damage multiplier by 0.25 once, regardless of matching piece count.']]
table(headers,rows,[.85,1.7,4.55])
p('The same amethyst-shard smithing pattern applies. Matching armor powers do not stack by piece count, but different powers can coexist. Only worn equipment supplies armor benefits. Arrows retain the power of the bow used at launch even when the held item changes.')

chapter('Regional creatures')
p('Ecosystems unlock after the first world dragon victory. They use native population caps and habitat-specific surface, underground, lighting or water checks. Hostiles do not naturally appear in Peaceful. Existing terrain can gain populations without regeneration. The mount habitats are independent of this progression gate.')
h('Forests')
use_table('| Creature / entity ID | Behavior and rewards |',[1.6,5.5])
p('Wooden creatures take extra damage from axes and fire. Barkling herb shearing is renewable, with a five-minute regrowth timer saved on the creature. Harvesting a calm Barkling does not anger it. The Ancient Ent provides an alternate Tree Spirit summon, not a boss heart.')
h('Taigas')
use_table('| Creature / entity ID | Behavior and player-kill drops |',[1.6,5.5])
p('Wolf packs hunt players and suitable prey. Pack speed bonuses do not multiply with every wolf. An Alpha has a bounded helper call and a warned pounce; called wolves expire and have no loot or XP. Alpha Fang remains a boss-only reward, while Alpha Dire Wolf grants the Moonlit Fang summon.')
_,regions=source_table('| Region | Creatures and discovery route |')
# Regional Catalog: Keep every registered creature visible, with source-derived health and temperament.
regional_source = (ROOT / 'src/main/java/darkspawn/black/ecosystem/RegionalKind.java').read_text(encoding='utf-8')
regional_stats = {name: (region, temper.lower(), health) for name, region, temper, health in re.findall(
    r'^\s*(\w+)\(BossProfile\.(\w+), Form\.\w+, Move\.\w+, Behavior\.\w+, Temper\.(\w+), (\d+),', regional_source, re.M)}
regional_notes = dict(line.split('|', 1) for line in '''ROTTED_ZOMBIE|Armored melee undead. Drops rotten flesh.
BRUTE_ZOMBIE|Warns before a ground smash. Opens and closes wooden doors without breaking them. Drops rotten flesh.
LEAPING_ZOMBIE|Fast undead with a warned pounce. Drops rotten flesh.
MUTANT_HUSKLING|Throws debris at range. Drops rotten flesh.
FAILED_MUTANT|Rare herald with a warned smash. Drops Mutated Flesh.
BONE_RAPTOR|Pack-spawning predator with a warned pounce. Drops a bone.
FOSSIL_SCORPION|Hides while idle; bites apply three seconds of poison. Drops a bone.
BONE_VULTURE|Flies, collects loose bones, and attacks players at half health or lower. Returns its carried item on death; ordinary drop is a bone.
FOSSILIZED_HUSK|Slow, heavily armored melee undead. Drops a bone.
TYRANT_SKULL|Stationary discovery on sand or terracotta. Right-click with a pickaxe for one Ancient Tyrant Fossil, costing one durability. Ordinary attacks cannot excavate it.
STORM_FINCH|Small passive flier with visible static during thunderstorms. Drops a feather.
SHOCKTALON|Warned diving attack briefly slows its victim. Drops a feather.
STORMSTRIDER|Fast neutral ground creature with visible static during thunderstorms. Drops a feather.
THUNDER_ROC|Rare flying herald fires lightning projectiles and knocks nearby targets back with wing pulses. Drops Charged Feather.
TREE_VIPER|Climbs walls, hides while idle, and applies three seconds of poison on a bite. Drops a vine.
CONSTRICTOR|Warns before attacking a fixed constriction area and strongly slowing victims. Drops a vine.
JUNGLE_STALKER|Hides until prey is close. Its cooldown-limited ambush hit deals extra damage and applies weakness. Drops a vine.
BROOD_SERPENT|Rare large serpent with a warned constriction attack. Drops Titan Scale.
BOGLING|Breathes underwater and scavenges loose food, spider eyes, or sugar; retreats with its item. Returns carried property on death. Ordinary drop is a spider eye.
HEXED_FROG|Melee hits apply either weakness or slowness for three seconds. Drops a spider eye.
SWAMP_WISP|Passive flying guide. Follow its particle-marked route to earn glow berries once. Ordinary player-kill drop is a spider eye.
HEXBOUND|Uses warned ranged hex projectiles. Drops a spider eye.
BABAS_FAMILIAR|Rare passive guide. Follow it to its marked destination for one Witch Token; killing it gives no token.
STONEBACK_GOAT|Neutral armored charger. Takes half damage from projectiles before other defenses. Drops flint.
CRAG_CRAWLER|Hides until prey approaches, then uses a stronger ambush bite with weakness on a cooldown. Drops flint.
STONEBORN|Slow armored rock creature that throws debris. Drops flint.
TITAN_SPAWN|Rare heavily armored herald throws debris and slams at close range. Drops Titan Stone.
FROSTLING|Hides while idle and fires frost projectiles. Drops a snowball.
FROZEN_HUSK|Armored melee undead with slowing hits. Fire deals double damage and briefly activates its thaw state. Drops a snowball.
ICEFANG|Pack-spawning predator whose bites apply three seconds of slowness. Drops a snowball.
YOUNG_WYRM|Rare herald uses ranged frost attacks and retreats below one-quarter health. Drops Wyrm Scale.
ABYSSAL_FISH|Hostile swimmer with melee bites; natural spawning requires dark water. Drops a prismarine shard.
GIANT_CRAB|Neutral armored seabed creature with a warned smash. Attacks from behind deal double damage before other defenses. Drops a prismarine shard.
SIREN|Warns before an attack that pulls victims toward it. Drops a prismarine shard.
LEVIATHAN_SPAWN|Rare deep-water herald with a warned grab, nearby boat damage, and a retreat into deeper water. Drops Abyssal Scale.
CAVE_SKITTERER|Small wall-climbing melee creature that can spawn in groups. Drops string.
WEB_SPITTER|Climbs and fires web projectiles, leaving temporary slowing patches without replacing blocks. Drops string.
BROOD_CARRIER|Armored climber releases a bounded group of temporary skitterer helpers on death. Drops string; helpers have no loot.
TUNNEL_WIDOW|Rare climber with a warned slowing pounce and a climbing retreat below half health. Drops Brood Fang.
SHADOW_DRONE|Warned close-range burst applies darkness and consumes the drone. Defeat it before the burst for a sculk vein.
SHADOW_STALKER|Stops moving when its target looks at it, then resumes stalking when unobserved. Drops a sculk vein.
SHADOW_SPITTER|Fires acid projectiles and leaves temporary darkness patches. Drops a sculk vein.
SHADOW_GUARDIAN|Armored guardian maintains a local home area and attacks with a warned shockwave. Drops a sculk vein.
SHADOW_LARVA|Natural larvae can mature into Shadow Drones after about one minute; temporary helpers cannot mature. Drops a sculk vein.
SHADOW_EGG|Stationary neutral egg alerts nearby regional allies when hurt and releases bounded temporary larvae on death. Drops a sculk vein.
SHADOW_PRAETORIAN|Rare armored herald uses warned directional strikes and darkness. Drops Shadow Membrane; Royal Shadow Gland remains boss-only.
SPORELING|Neutral hive creature alerts allies when attacked and grants nearby regional allies brief speed during combat. Drops a brown mushroom.
MYCELIUM_CRAWLER|Hides until prey is close, then uses a stronger ambush bite with weakness on a cooldown. Drops a brown mushroom.
INFECTED_MOOSHROOM|Fires spores and poisons nearby attackers with a cooldown-limited defensive spore response. Drops a brown mushroom.
SPOREWALKER|Ranged spore attacker. Drops a brown mushroom.
MYCELIAL_GUARDIAN|Rare neutral herald heals one health every two seconds on mycelium and creates temporary spore patches in combat. Drops Mycelial Heart Fragment.
CRIMSON_RAVAGER|Crimson Forest charger with a warned attack. Drops crimson fungus.
FUNGAL_IMP|Crimson Forest melee creature grants nearby regional allies brief speed during combat. Drops crimson fungus.
BLOODROOT|Stationary Crimson Forest predator uses warned close-range snares. Drops crimson fungus.
WARPED_STALKER|Warped Forest creature becomes hostile when stared at and can blink during combat. Drops warped fungus.
WARPED_WISP|Warped Forest flying spore attacker with short combat blinks. Drops warped fungus.
RIFTLING|Neutral Warped Forest flier with short combat blinks. Drops warped fungus.
ASH_GHOUL|Nether Wastes and Basalt Deltas melee creature ignites targets for two seconds. Drops magma cream.
MAGMA_BRUTE|Armored Wastes and Deltas creature uses warned fiery attacks and temporary magma patches. Drops magma cream.
INFERNAL_HOGLIN|Wastes and Deltas charger with a warned attack. Drops magma cream.
CRIMSON_SPAWN|Rare Crimson Forest herald with a warned charge. Drops Crimson Organ.
WARPED_SPAWN|Rare Warped Forest herald uses spore projectiles and short combat blinks. Drops Warped Organ.
INFERNAL_SPAWN|Rare Wastes and Deltas herald uses fiery attacks and temporary magma patches. Drops Infernal Organ.
LOST_SOUL|Flying spirit briefly passes obstacles during pursuit. Contact attacks temporarily reduce maximum health by half a heart. Drops a bone.
BONEWALKER|Melee undead reconstructs once at half health after ordinary lethal damage. That used revival persists across reloads. Drops a bone on final defeat.
SOULFLAME_SKULL|Flying ranged soul-bolt attacker. Drops a bone.
TORMENTED|Stationary predator uses warned close-range snares. Drops a bone.
SOUL_KEEPER|Rare herald fires soul bolts that temporarily reduce maximum health by one heart and calls bounded Lost Soul helpers. Drops Captured Soul.
VOIDLING|Small hostile flier with short combat blinks; natural groups can contain several creatures. Drops chorus fruit.
END_GRAZER|Passive creature heals near chorus without consuming the plants and blinks away when injured. Drops chorus fruit.
SHARDLING|Neutral flier can blink to dodge a projectile when its cooldown is ready. Drops chorus fruit.
VOID_RAY|Large neutral flying melee creature. Drops chorus fruit.
WATCHER|Flying ranged observer retreats when prey is close. Drops chorus fruit.
VOID_SENTINEL|Rare herald uses void projectiles, combat blinks, rotating defensive fragments, and bounded Voidling helpers. Strike between its shield fragments. Drops Watcher Fragment.'''.splitlines())
assert set(regional_notes) == set(regional_stats), 'Regional guide must describe every registered kind exactly once'
region_keys = ['MUTANT_ZOMBIE', 'FOSSIL_TYRANT', 'THUNDER_BIRD', 'TITAN_BOA', 'BABA_YAGA',
               'MOUNTAIN_TITAN', 'ICE_WYRM', 'KRAKEN', 'CAVE_CRAWLER', 'SHADOW_CREEPER_QUEEN',
               'MYCELIAL_SOVEREIGN', 'NETHERBORN', 'SOULBOUND_COLOSSUS', 'VOID_EYE']
assert len(regions) == len(region_keys), 'Regional source table order changed'
p('The following catalog describes every additional registered regional creature. Health is measured in points; two points equal one heart. Ordinary death drops listed here require a player kill and normal mob loot rules. Summoned helpers provide no ordinary loot. Rare heralds supply discovery materials, not boss hearts.')
for (region,body), region_key in zip(regions, region_keys):
    h(region.replace(' / ',' and '))
    p(body).paragraph_format.keep_with_next = True
    rows = []
    for name, (key, temper, health) in regional_stats.items():
        if key != region_key:
            continue
        entity_id = name.lower()
        display = LANG.get('entity.darkspawn.' + entity_id, entity_id.replace('_', ' ').title())
        rows.append([f'{display}\n{entity_id}', f'{health} HP; {temper}. {regional_notes[name]}'])
    table(['Creature and entity ID', 'Health behavior and rewards'], rows, [1.75, 5.35])
h('Discoveries and population limits')
p("Regional discovery materials combine with dragon's breath to produce their existing boss summons. Thunder Bird needs a Stormforged Feather; Netherborn needs all three regional organs. Discoveries do not award boss hearts or replace ritual weather, space and biome requirements.")
p('Special attacks warn before executing at fixed locations. Temporary web, darkness, magma and spore areas do not replace player blocks. Helper producers have bounded lifetime calls, local caps, no helper loot and short expiration. Natural Shadow Larvae can mature once; summoned helpers cannot reproduce. Teleports require loaded, collision-free, supported destinations.')
h('Endborn')
p('Endborn are separate hostile End creatures with 60 health, melee attacks, teleports and a projectile-dodge cooldown. They naturally appear after the dragon victory in End highlands, midlands and barrens. Player-killed natural Endborn drop a shard and pearl. Void Eye helpers do not provide these drops. Endborn Shards help craft the Watcher Lens.')

chapter('Cooking and agriculture')
p('The Cooking Station combines one to four ingredient types in any order. It previews the result before taking ingredients, works instantly, and needs no fuel. Taking an output consumes one of each listed ingredient; bulk crafting stops when output space runs out. Buckets and bottles return appropriately. Inputs belong to the open player menu and return when it closes; the station does not store shared ingredients.')
p('Craft the station from four iron ingots, a campfire and three planks. Meals stack to sixteen and can be eaten at full hunger. Only one meal buff is active: another meal replaces it. Ordinary foods and cheese do not replace the buff, and independently applied potions keep their own timers. Meal timers save, pause offline, and clear on death or milk.')
h('Crops and ingredients')
use_table('| Ingredient | Raw hunger / saturation |',[1.05,1.1,2.4,2.55])
p('The six ordinary crops grow on farmland through eight ages, using normal light, hydration and bone-meal rules. Mature crops yield produce and seeds; Fortune increases the extra-seed roll. Farmers can plant and harvest them, chickens accept the seeds, and produce and seeds can be composted. Farmers may offer four seeds for two emeralds through their randomized novice trade pool.')
p('Frost Garlic, Jungle Pepper and Marsh Rice add three regional crops. Ordinary garlic, pepper or rice has a 25% mutation chance at the final natural growth step in its matching frozen, jungle or swamp region. Bone meal does not trigger mutation. Once obtained, regional seeds can grow on suitable farmland elsewhere; regional produce converts back to its ordinary ingredient.')
h('Cookbook discoveries and feasts')
p('Craft the Cookbook from a book and Wild Herbs. It lists categorized recipes, ingredients, effects, collected meals and hints for locked dishes. Right-click recipe scrolls to learn them permanently; duplicate scrolls are preserved.')
table(['Discovery','Current sources'],[
('Shadow Stew','Ancient City and mineshaft chest loot.'),('Infernal Roast','Bastion treasure chest loot.'),('Captain\'s Chowder','Shipwreck supplies, fishing treasure and higher-level Chef trades.'),('Frost Garlic Soup','Igloo chest loot.'),('Jungle Pepper Skewer','Jungle temple chest loot.'),('Marsh Rice Bowl','Dungeon chest loot.'),('Bog Gumbo','Journeyman Chef.'),('Hero Feast','Chef scroll or collecting twenty distinct Darkspawn meals.')],[1.6,5.5])
p('Loot-table scroll additions use a 50% roll when the applicable loot is generated. Existing opened inventories are not rewritten. A Cooking Station can serve as a Chef workstation for an unemployed adult villager. Chefs use five native trade levels, stock and restocking; they cannot take another player\'s station inputs.')
p('Hunter\'s, Woodland, Ocean, Nether and Hero Feasts, plus Ender Banquet, each provide six servings. Place a feast and empty-hand right-click while hungry. A serving gives eight hunger points and a five-minute effect. Untouched feasts can be recovered; breaking a partly eaten feast loses the remaining servings.')
h('Meal effects and variety')
use_table('| Meal identity | Benefit |',[1.75,5.35])
p('The recent-food window remembers eight eating events, including vanilla foods. Three distinct foods add 5% bonus saturation, five add 10%, and eight also heal half a heart every five seconds while at full hunger. Repeated food still gives normal nutrition. Recent variety resets on death; learned recipes and lifetime meal collection persist.')
p('The appendix lists all 59 current Cooking Station outputs and ingredients. Each boss premium dish uses its own boss ingredient with rice, butter and Wild Herbs. Decorative trophies are separate placeable rewards, not ingredients consumed by the meal.')

chapter('Armor and climate')
h('Standard armor styles')
p('Iron and diamond keep their vanilla item identities and core protection. Fifteen styles replace their appearance and grant a bonus when all four worn pieces share a style. Iron and diamond may mix. Inactive loadouts and incomplete or mixed-style sets grant no style bonus.')
use_table('| Style | Full-set bonus |',[1.25,5.85])
p('Apply a style at a smithing table with paper, an iron or diamond piece, and the style addition. Styling preserves names, enchantments, damage, trims and boss powers. Applying the current style wastes nothing. Restore Classic with an iron nugget. Retired Arbalist, Royal Guard and Brawler data loads as Ranger, Paladin and Berserker.')
use_table('| Appearance | Addition |',[2.1,5.0])
figure('art/armor-styles/final-roster-preview.png','Standard armor asset preview, showing the later Ranger, Paladin and Shadowblade designs. Equipment is shown in front and rear views.',max_height=4.5)
h('Legendary armor').paragraph_format.page_break_before = True
p('Five four-piece legendary families upgrade netherite armor and keep its base protection and repair rules. Each piece requires a netherite upgrade template and a core crafted from two specified boss essences plus a netherite ingot. A full set requires four cores and templates. Only a complete worn matching set grants its legendary power.')
use_table('| Set | Essences needed for each core |',[1.05,1.8,4.25])
p('Phoenix and Soul Reaper cooldowns persist through equipment swaps, reconnects, restarts and respawns, using world ticks. Phoenix does not protect against the void or grant lasting invulnerability. Legendary smithing preserves item data, and standard paper-style recipes cannot convert legendary pieces.')
figure('art/armor-styles/legendary-preview.png','Legendary armor asset preview. Front and rear pairs: Dragonlord and Phoenix, Voidwalker and Titan, then Soul Reaper.',max_height=3.8)
h('Expedition equipment').paragraph_format.page_break_before = True
p('Five expedition families convert matching iron armor using leather in the template slot and a family-specific addition. Worn pieces supply benefits; stored sets do not. Families can mix, but full-suit effects require their stated pieces. Helmets, pouches, tanks and canteens are modeled details; they do not add cargo, refill tanks or dynamic light systems.')
use_table('| Set | Defense: helmet / chest / legs / boots |',[1.2,1.0,4.9])
use_table('| Output | Addition |',[2.6,4.5])
figure('art/expedition-armor/preview.png','Expedition armor asset preview. Front and rear pairs: Mountaineering and Spelunking, Diver and Desert Wanderer, then Ashwalker.',max_height=3.8)
h('Temperature and preparation')
p('Survival and Adventure players track heat and cold exposure, shown by a HUD meter with direction and wetness. Biome, exposed altitude, night, rain, water, cover and nearby heat sources affect conditions. This is a gameplay exposure scale, not degrees Celsius. Creative and Spectator are exempt.')
p('Exposure changes once per second: one point toward worsening conditions or two points toward recovery. At 35 points the HUD warns; at 60, movement slows by 10% and mining by 15%; at 85, climate damage is one health point every five seconds. Shelter and appropriate equipment gradually improve exposure rather than instantly clearing it.')
p('Mountaineering protects against cold; Desert Wanderer and Ashwalker protect against heat; Diver adds cold protection while wet. Warmth and Cooling I-II add protection on worn armor and are mutually exclusive on one piece. Warming meals help against cold, and Fire Resistance helps against heat without granting complete climate immunity. Deep cover moderates Overworld climate, while Nether heat persists under cover.')
p('Exposure and wetness save across reconnects and living dimension travel; offline exposure pauses, and death starts comfortable. There is no thirst system. Spelunking night vision is personal vision, not emitted world light or mob-spawn suppression.')

chapter('Inventory and storage')
h('Personal capacity and sorting')
p('Press R during gameplay or use Sort inventory in the Survival inventory. Sorting merges compatible stacks and orders the main inventory plus extra rows by item ID. It preserves hotbar, armor, offhand and crafting slots. Cursor-held items, stale menu state and unsupported containers prevent unsafe requests.')
use_table('| Upgrade | Crafting ingredients | Total main slots |',[1.1,4.7,1.3])
p('Apply leather, iron and diamond upgrades in order by holding and using them. Each adds nine slots, reaching 54 main slots plus the separate nine-slot hotbar. Capacity is permanent per player; stored items follow normal death and keepInventory rules. Extra rows can receive overflow pickups and provide recipe-book ingredients.')
h('Armor loadouts')
use_table('| Upgrade | Crafting ingredients | Result |',[1.75,3.65,1.7])
p('Iron unlocks set two and diamond unlocks set three. The inventory shows the active, locked or switching state. To populate a set, switch to it and equip the normal armor slots. Inactive sets grant no stats and use no main-inventory space.')
p('Switching takes two seconds while grounded with the inventory open. Damage, movement into the air, item use, inventory changes, dimension travel or disconnects cancel it. Riding, gliding, cursor-held items, recent damage and Curse of Binding can prevent a switch. Saved items keep all their data. Inactive armor follows normal death drops, vanishing and keepInventory rules.')
h('Directed storage routes')
p('Build a route from Dump Chest through Sorting Tubes and Sorting Chests to a terminal Bank Chest. Gold arrows indicate output direction. Empty-hand use rotates tubes; sneak-use rotates chests. Filters record item types without consuming the sample.')
use_table('| Block | Purpose | Crafting ingredients |',[1.1,3.7,2.3])
p('A Dump attempts up to 64 items from one slot every five ticks, rotating slots to avoid starvation. Matching Sorting Chests receive what fits, then remaining items continue to later chests and the Bank. Routing checks at most 128 loaded blocks, detects cycles and never force-loads chunks. Unaccepted items stay at their source.')
p('Sorting Chests have 27 real slots and three ghost filters. Empty filters accept no automatic stock; manual deposits remain unrestricted. Hoppers can access Dump and Sorting Chests, with filter enforcement on insertion. Banks use their menus and tubes instead of direct hopper access.')
h('Bank storage and recovery')
p('Banks start at 256 distinct item variants with up to 1,000,000 items per variant. Named, damaged or enchanted variants remain separate. Upgrades costing one, two and four diamond blocks raise capacity to 512, 1,024 and 2,048 variants. Banks and shulker boxes cannot be nested inside banks; bank items cannot enter bundles or shulker boxes.')
p('Search by item name or ID, sort by name or count, page results, withdraw one item or a stack, and bulk-deposit eligible inventory. Re-sort sends stored stock back toward eligible Sorting Chests. The separate Trash view is recoverable until Empty trash is confirmed with a second click within five seconds. Manual deposits never automatically enter Trash.')
p('Breaking a Bank preserves its inventory, filters, Trash and upgrades in one Bank item. Dump and Sorting Chests drop their contents normally. Dropped banks can be lost like any other dropped valuable item; pick them up promptly.')

chapter('Bonfires and exploration structures')
h('Barracks Bonfire')
p('Craft the bonfire from two iron ingots, two copper ingots, a campfire and three stone bricks. Place it over the center full copper block of a three-by-three foundation, with stone bricks in the other eight positions. Full copper blocks support wax and oxidation variants; cut copper is not valid.')
p('Empty-hand right-click toggles power. Sneak-use cycles horizontal radius through 25, 50 and 75 blocks. Coverage extends ten blocks below and twenty layers above the core. Blue flames mean activation is still clearing hostiles or placing initial lights; orange means active. No fuel is consumed.')
p('The ward blocks ordinary hostile spawning and entry, including spawners, eggs and raids inside its area. Entire collision boxes are checked. Existing hostiles are moved to safe loaded ground outside all wards; if relocation cannot succeed, they are paused and activation remains blue. Bosses, Withers and Ender Dragons are exceptions. Projectiles and explosions from outside are not blanket-blocked.')
p('The prototype places replaceable light sources on a sparse grid, providing real block light without replacing solid blocks or water. Walls still block light and dark pockets can remain. Core and foundation must be loaded; the ward never force-loads chunks. Breaking or disabling it removes its own unneeded lights when those chunks tick, while overlapping wards retain shared lights.')
h('Boss discovery landmarks')
p('Overgrown Druid Shrines generate in forests and Abandoned Hunter Camps in taigas. Each has three variants and four rotations, a ritual clue book and modest finite loot. Shrine footprints are 31 by 31 blocks; camps are 25 by 25. Candidate placement uses 40-chunk spacing and 16-chunk separation, with terrain rejection, so this is not a guaranteed travel distance.')
p('All sites can be explored before defeating the dragon. They do not grant hearts, premium boss ingredients or complete summons. Boss rituals still require progression and valid conditions. These structures appear only in newly generated terrain; existing explored chunks and opened inventories are not rewritten. All fourteen additional boss families now have regional sites with three variants each. The catalog below lists their native structure IDs.')
use_table('| Region | Landmark ID | Boss |',[1.6,2.65,2.85])
p('New regional sites occupy 25 by 25 blocks. Surface and End sites require supported gentle terrain; cavern sites require natural boss headroom, a retained roof and an opening into the cavern. Wrecks require deep open water over a supported seabed. The Kraken ritual still takes place at the surface. Native exclusion zones separate farm excavation candidates from villages and hive candidates from Ancient Cities. Conservative terrain checks can make sites scarce, particularly in Mushroom Fields.')
h('Research bunkers')
p('Research bunkers generate in new Overworld terrain and introduce Forsaken access. Their Quarantine Protocol book and research archive explain gate construction and fabrication. Each visitor can learn the schematic permanently; an archive is not consumed by the first player.')

chapter('OneBlock and Forsaken')
h('Shared OneBlock island')
p('Use the OneBlock button in the inventory to enter the shared island for free. Inside, Return goes back. Travel takes three seconds standing still on solid ground; damage cancels it and adds a combat delay. Dismount and detach carried or leashed pets first, close other containers and clear the cursor and crafting grid.')
use_table('| Shared phase | Breaks to unlock the next phase |',[1.1,1.15,4.85])
p('Everyone shares one persistent island and phase. The resource at 0, 79, 0 regenerates above a protected bedrock anchor. Its represented vanilla block controls mining speed, harvest requirements, drops and experience. Counts are per phase. Meeting a threshold unlocks manual Next phase rather than forcing the group onward.')
p('A supply chest is attempted every 25 successful breaks and a vanilla creature every 50, only when safe supported space is available. Earlier essentials remain in later pools. Creative breaking does not advance phases. No custom Darkspawn boss events or distant islands are included.')
p('Possessions are separate: inventory, hotbar, equipment, extra rows, stored armor and Ender Chest contents switch to saved OneBlock versions and back on return. Character upgrades, XP, health, hunger and effects remain shared. There are no normal item imports or exports. Death respawns on the island and follows keepInventory; void losses are real, while outside possessions remain saved.')
h('Forsaken city expeditions')
p('Learn the bunker schematic, repair the five-by-five obsidian gate frame with a three-by-three opening, and fabricate a reusable Containment Key. After world dragon victory, use the key on the controller to activate the gate. Empty-hand use enters. Researched players can fabricate personal gate parts, and everyone can use an already opened gate.')
p('The landing has an Extraction Terminal that returns each traveler near their own entrance. Equipment and salvage travel normally; Forsaken does not use OneBlock item separation. Gates have no fuel timer, but a damaged frame must be repaired and reactivated. Normal death and keepInventory rules apply.')
p('The first city slice lies south of the landing. Apartments and shops lead through quarantine boundaries to a hospital and research annex. Ten finite scavenging caches support ground, upper-floor, rooftop and tunnel exploration. A rooftop crossing offers another route, and the main retreat street remains clear. Cities repeat every 512 blocks along both axes in fresh terrain.')
p('Four infected roles now threaten the city: Wanderers charge, Crawlers flank, Spore Spitters attack from range and retreat when crowded, and slow armored Brutes appear under high pressure. Their attack windups give time to move away or take cover. Four permanent nests occupy the hospital and annex in new cities. Destroying a nest yields one Parasitic Core and suppresses unrelated city spawns for five minutes; other intact nests remain active.')
p('Contamination slows movement by 15% for at most 60 seconds, refreshed by later exposure. Honey, milk, successful extraction or death clears it. Each worn Quarantine Armor piece shortens new exposure by 15 seconds; a full set prevents new contamination. It does not cure an existing effect. The suit has iron-tier defense, a sealed hood, respirator filters and a rear scrubber.')
p('Combat, explosions and breaking blocks raise one shared pressure value per city. Quiet time reduces it. Pressure is capped at 100, with warnings at 25, extra activity at 50 and a short surge at 75. A surge lasts at most 30 seconds with two minutes between starts. Spawning is limited to 12 nearby infected per city, 64 loaded infected across Forsaken and eight total spawn attempts per second. Peaceful and disabled mob spawning are respected.')
figure(ROOT/'art/forsaken/content-preview.png','Forsaken enemies and quarantine equipment. Authored geometry preview, not an in-game screenshot.',width=6.4,max_height=2.7)
h('Forsaken resources and equipment')
table(['Resource','Where to find it','Implemented purpose'],[
('Contaminated Alloy','Salvage caches, 1 to 3 guaranteed','Modules, scanner and beacons'),
('Biofiber','Player-killed infected; medical caches','Modules and armor smithing'),
('Parasitic Core','One per destroyed nest','Modules, scanner, blade and beacon fuel'),
('Research Data','Annex and rooftop caches, 1 to 3','Modules, scanner and beacons'),
('Quarantine Module','20% extra roll in medical, research and roof caches; craftable','Armor and blade upgrades; beacons')],[1.4,3.2,2.5])
p('Craft a Quarantine Module with 4 Biofiber, 2 Alloy, 2 Research Data and 1 Core. At a smithing table, use a module as the template, an iron armor piece as the base and Biofiber as the addition. This preserves the gear’s existing damage, name, enchantments and other native components. Repeat for helmet, chestplate, leggings and boots.')
p('Craft the Bio-Scanner from 1 Alloy, 2 Research Data, 1 Core and 1 Redstone. Use it in Forsaken to scan loaded terrain within 32 blocks. Sneak-use cycles nests, resources and caches, and local infestation; all modes share a five-second cooldown.')
p('Existing cities and opened caches are preserved. Find a fresh city for nests; unopened caches use updated loot. City pressure survives restarts. Unloaded nests do not accumulate spawn waves.')

h('Forsaken expedition tools')
doc.paragraphs[-1].paragraph_format.page_break_before = True
p('These tools extend the city expedition loop with infected control, a temporary refuge from new spawns and information for scavenging. Existing enemies can still enter a beacon area.')
figure(ROOT/'art/forsaken/expedition-tools-preview.png','Authored textures and beacon front projection, not a gameplay screenshot.',width=6.4,max_height=2.0)
h('Pulse Blade')
p('At a smithing table, combine a Quarantine Module template, an iron sword and a Parasitic Core. The upgrade preserves names, enchantments and damage. Ordinary melee deals 25 percent extra damage to infected. Hold Use for 1.5 seconds, then release a four-block pulse: 10 damage to visible, non-allied infected, an eight-second cooldown and eight durability per pulse. Early release is free. Players, pets and blocks are unaffected by the pulse. Charge sparks, a ready chime and a visible ring mark its timing.')
h('Containment Beacon')
p('Craft one device from 4 Alloy, 2 Modules, 2 Redstone and 1 Research Data. Place it in Forsaken and use one Core to power it for five minutes. It suppresses new managed infected spawns within a three-dimensional 16-block radius. Overlapping devices do not stack or kill existing enemies. Use an empty hand to check remaining time. Active top-ups are refused; breaking it returns the device and loses the remaining fuel. Fuel deadlines survive saves, and expiry follows server world time.')
h('Scanner modes')
p('Sneak-use cycles Nests, Resources and caches, and Local infestation. Nests reports the closest intact nest; Resources reports the closest dropped exclusive material or unopened Forsaken cache. Cache signals never open containers or generate loot. Local infestation reports the current city pressure, district and remaining surge time. Scans stay in loaded terrain, share a five-second cooldown, and cannot reveal distant unloaded cities.')

chapter('Darkspawn Journal')
p('Open the free Journal button in your inventory or expanded Creative storage screen. It brings creatures, mounts, bosses, cooking, Forsaken and landmark discoveries into one searchable reference. The world keeps running while you read. Use Refresh after exploring to update the server snapshot.')
figure(ROOT/'art/journal/discovery-preview.png','Journal portraits before and after discovery, derived from the existing creature models.',width=6.4,max_height=2.7)
h('Discovery belongs to the player')
p('Observe a living creature within 24 blocks with a clear line of sight, or enter a native landmark, to discover it. Checks run every two seconds. Undiscovered species show silhouettes and habitat hints; known entries reveal names, behavior, rewards and controls. Discovery survives death, reconnects and OneBlock item-profile changes. Spectators do not gain observations. Previously visited landmarks must be revisited if no record exists.')
h('Existing progress carries over')
p('Cooking entries read current Cooking Station recipes and your existing cookbook discoveries. Consumed boss hearts and Minecraft credited-kill statistics are separate: receiving a heart never invents a kill. Historical shared participation cannot be reconstructed. Visiting a boss landmark reveals its ritual clues. Existing Forsaken research unlocks expedition guidance; infected portraits unlock through observation or a credited kill.')
h('Browse and search')
p('Choose Guide, Wildlife, Mounts, Bosses, Cooking, Forsaken or Sites. Search filters the displayed information in that section. The left arrows page through entries; Previous and Next page through long descriptions. Section counters show known and total entries. The catalog covers 82 regional creatures, 14 mounts, 16 bosses, four infected and 17 landmark families, plus guidance and current cooking recipes.')
p('This first pass is an English field reference with static model illustrations. It never opens caches, loads distant terrain, awards rewards or moves possessions. Discovery coordinates, maps, structure compasses and archaeology remain future exploration work. Actual GUI scaling, narration and multiplayer readability remain live acceptance checks.')

chapter('Presentation and prototype boundaries')
p('The creature pipeline contains 115 regional and boss GeckoLib assets plus four Forsaken infected rigs, separate from the 14 Java-rendered mount types. Editable Blockbench sources, textures and runtime models accompany the content. Mounts now have species-specific materials and animation, and the Nyan-inspired cat emits a six-color rear trail when ridden.')
p('Regional creatures have native-state-driven animations and custom sound work. Boss encounters have phase-based music, with volume reduced during attack warnings. Some impacts, utility icons and effects still use vanilla assets. The mod does not include the Nyan Cat song.')
p('Cooking labels truncate with full-name tooltips when necessary. Temperature text wraps to fit the viewport. Armor inventory textures and trim borders are registered on the same atlas. These presentation changes do not change equipment recipes or damage rules.')
h('Explicitly unfinished or planned')
table(['Area','Not currently implemented'],[
('Mount combat','Autonomous combat behavior and further species abilities.'),
('Transport concepts','Magic carpet, cloud, airplane and mech systems.'),
('Forsaken','Procedural city expansion and long-expedition balance review.'),
('World exploration','Expanded archaeology, discovery maps, rematches and Nexus concepts. All sixteen boss landmark families are implemented; live terrain and rarity acceptance remains.'),
('Character systems','Proficiency and skill-gated crafting; thirst.'),
('OneBlock','Extra islands, custom boss events and normal export/import of outside items.'),
('Equipment details','Diver tank refill, boot anchoring and storage inside modeled expedition pouches.')],[1.45,5.65])
p('The implementations above are present in development source. Alpha balancing, final animation and camera fit, encounter frequency and multiplayer behavior can still change. This guide is not a claim that every feature has completed live acceptance testing.')
h('Source record')
p('Prepared from the current Darkspawn registrations and implementations, the player README, mount and bonfire notes, OneBlock and Forsaken implementations, armor catalogs, localized names and Cooking Station recipe resources. Source snapshots and document generation are dated October 13 2026. The release version is unchanged; no new playable JAR is implied by this guide.')

chapter('Cooking Station recipe catalog')
p('Every ingredient below costs one item per craft. Alternatives separated by “or” satisfy one ingredient slot, not multiple slots. Recipe discoveries still apply. The current catalog has 59 recipe outputs; the Cookbook shows recipe-specific effect details in-game.')
cuisine=(ROOT/'src/main/java/darkspawn/black/cooking/Cuisine.java').read_text()
discoveries={m[0]:m[1] for m in re.findall(r'addMeal\("([^"]+)"[^\n]*?, "([^"]*)", (?:true|false)\);',cuisine)}
for name,discovery in re.findall(r'addFeast\("([^"]+)"[^\n]*?, "([^"]*)"\);',cuisine):discoveries[name]=discovery
def ingredient(value):
    if isinstance(value,list):return ' or '.join(ingredient(v) for v in value)
    if isinstance(value,str):return item_name(value) if not value.startswith('#') else value[1:].split(':')[-1].replace('_',' ')
    return ingredient(value.get('item',value.get('items',value.get('tag','Unknown ingredient'))))
recipes=[]
for f in (ROOT/'src/main/resources/data/darkspawn/recipe').glob('*.json'):
    r=json.loads(f.read_text())
    if r.get('type')!='darkspawn:cooking':continue
    result=r['result'];id=result['id'];name=item_name(id)
    if result.get('count',1)!=1:name+=' x'+str(result['count'])
    gate=discoveries.get(id.split(':')[-1],'') or 'Available'
    recipes.append([name,' + '.join(ingredient(i) for i in r['ingredients']),gate])
recipes.sort(key=lambda r:r[0])
table(['Output','Ingredients','Discovery requirement'],recipes,[1.65,3.85,1.6])
path=OUT/'Darkspawn Implementation Guide.docx'
doc.save(path)
print(path)
print(f'Recipe entries: {len(recipes)}; embedded images: {len(doc.inline_shapes)}')
